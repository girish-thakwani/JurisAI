import re
from playwright.sync_api import Playwright, sync_playwright, expect


def pause_between(page, lo=2000, hi=5000):
    import random
    page.wait_for_timeout(random.randint(lo, hi))




class Row(dict):
    # One data row. row["phone"] reads a column; row.set("status", "Completed") writes it back to the source.
    _write = None

    def set(self, col, value):
        if isinstance(value, str):
            value = " ".join(value.split())  # one line, so sheet cells stay tidy
        self[col] = value
        if self._write:
            try:
                self._write(self, col, value)
            except Exception as e:
                print(f"could not write {col} back to the data source: {e}")


class Rows(list):
    columns = []


def load_rows(source, header=True, creds="creds.json", sheet=None):
    # Read rows from a CSV / Excel file or a Google Sheet URL. Values stay text (keeps +91 and leading zeros).
    import re
    if re.match(r"https?://docs.google.com/spreadsheets", source):
        import gspread
        book = gspread.service_account(filename=creds).open_by_url(source)
        ws = book.worksheet(sheet) if sheet else book.sheet1
        grid = [list(r) for r in ws.get_all_values()]

        def save(r, c, v):
            ws.update_cell(r + 1, c + 1, v)
    elif source.lower().endswith((".xlsx", ".xlsm")):
        import openpyxl
        wb = openpyxl.load_workbook(source)
        ws = wb[sheet] if sheet else wb.active

        def text(c):
            if c is None:
                return ""
            return str(int(c)) if isinstance(c, float) and c.is_integer() else str(c)
        grid = [[text(c) for c in r] for r in ws.iter_rows(values_only=True)]

        def save(r, c, v):
            ws.cell(row=r + 1, column=c + 1, value=v)
            wb.save(source)
    else:
        import csv
        with open(source, newline="", encoding="utf-8-sig") as f:
            grid = [list(r) for r in csv.reader(f)]

        def save(r, c, v):
            with open(source, "w", newline="", encoding="utf-8") as f:
                csv.writer(f).writerows(grid)

    def col_letter(i):
        name, i = "", i + 1
        while i:
            i, r = divmod(i - 1, 26)
            name = chr(65 + r) + name
        return name

    width = max((len(r) for r in grid), default=0)
    if header and grid:
        keys, body, first = [str(h).strip() for h in grid[0]], grid[1:], 1
    else:
        keys, body, first = [col_letter(i) for i in range(width)], grid, 0
    rows = Rows()
    rows.columns = list(keys)

    def write(row, col, value):
        if col not in keys:
            if not header and re.fullmatch(r"[A-Z]+", col):
                while col not in keys:
                    keys.append(col_letter(len(keys)))
            else:
                keys.append(col)
                if header:
                    grid[0] += [""] * (len(keys) - 1 - len(grid[0]))
                    grid[0].append(col)
                    save(0, len(keys) - 1, col)
            rows.columns = list(keys)
        r, c = row._pos, keys.index(col)
        grid[r] += [""] * (c + 1 - len(grid[r]))
        grid[r][c] = value
        save(r, c, value)

    for i, cells in enumerate(body):
        cells = list(cells) + [""] * (len(keys) - len(cells))
        if not any(str(c).strip() for c in cells):
            continue
        row = Row(zip(keys, (str(c).strip() for c in cells)))
        row._pos, row._write = first + i, write
        rows.append(row)
    return rows




def run(playwright: Playwright) -> None:
    browser = playwright.chromium.launch(headless=False)
    context = browser.new_context()
    page = context.new_page()
    page.goto("https://web.whatsapp.com/")
    page.get_by_role("button", name="Continue").click()
    rows = load_rows("https://docs.google.com/spreadsheets/d/1CuR0wgOiO5-P1Pw6lMdoiVjjhWrpYco6b2U-hfMK7Mc/edit?gid=0#gid=0", sheet="Sheet1")
    for row in rows:
        if row.get("status") == "Completed":
            continue
        try:
            page.get_by_role("textbox", name="Search or start a new chat").click()
            page.get_by_role("textbox", name="Search or start a new chat").fill("note")
            page.get_by_test_id("conversation-compose-box-input").click()
            page.get_by_test_id("conversation-compose-box-input").fill("hello")
            row.set("status", "Completed")
        except Exception as e:
            row.set("status", f"Failed: {str(e)[:100]}")
        pause_between(page, 3, 4)
    page.close()

    # ---------------------
    context.close()
    browser.close()


with sync_playwright() as playwright:
    run(playwright)
